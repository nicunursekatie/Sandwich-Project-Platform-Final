"""Create the full corrected version of the supplied collection findings report.

Production collection-log figures were checked read-only on September 28, 2026.
The Word document and PDF are generated from the same data and narrative.
"""

from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

OUTPUT = Path("reports")
OUTPUT.mkdir(exist_ok=True)
DOCX = OUTPUT / "Peak_Sandwich_Collection_Weeks_Full_Corrected.docx"
PDF = OUTPUT / "Peak_Sandwich_Collection_Weeks_Full_Corrected.pdf"

annual = [
    ("2020*", "44,958"), ("2021", "281,395"), ("2022", "440,371"),
    ("2023", "473,409"), ("2024", "501,856"), ("2025", "526,083"),
]
months = [
    ("November 2022", "60,082"), ("January 2025", "59,022"),
    ("October 2024", "55,253"), ("November 2023", "54,849"),
]
locations = [
    ("Dunwoody/PTC", "90,011"),
    ("Alpharetta", "74,798"),
    ("East Cobb/Roswell", "69,520"),
    ("Intown (combined name variants)", "55,441"),
    ("Sandy Springs/Chastain (combined name variants)", "24,701"),
    ("Flowery Branch", "12,706"),
]
seasonal = [
    ("Jan", "48,273", "Jul", "36,283"),
    ("Feb", "43,493", "Aug", "32,849"),
    ("Mar", "39,256", "Sep", "38,368"),
    ("Apr", "39,830", "Oct", "43,714"),
    ("May", "35,727", "Nov", "49,900"),
    ("Jun", "40,949", "Dec", "36,791"),
]
peak_weeks = [
    ("Jan 15-21, 2025", "28,927", "7,361", "21,566"),
    ("Nov 15-21, 2023", "27,810", "8,136", "19,674"),
    ("Jan 11-17, 2023", "27,132", "5,134", "21,998"),
    ("Jan 14-20, 2026", "24,717", "4,570", "20,147"),
    ("Sep 25-Oct 1, 2024", "24,664", "6,617", "18,047"),
]
weekly_medians = [
    ("2022", "50", "8,168"), ("2023", "50", "8,817"),
    ("2024", "50", "9,121"), ("2025", "51", "9,471"),
]

# Entries are headings, paragraphs, notes, or tables. This list is shared by
# both output formats so they cannot quietly diverge.
report = [
    ("title", "Key Findings: Sandwich Collection Trends and Peak Weeks"),
    ("subtitle", "An evidence-based review of recorded sandwich collections, April 2020-September 2026"),
    ("note", "Prepared September 28, 2026. Weekly rankings include completed weeks through September 22, 2026."),
    ("heading", "Key takeaways"),
    ("text", "The highest completed Wednesday-Tuesday collection week on record was January 15-21, 2025, with 28,927 sandwiches logged. The next highest was November 15-21, 2023, with 27,810. These are totals of individual and group sandwiches in current collection records."),
    ("text", "Collections increased across the last four full calendar years: 440,371 in 2022, 473,409 in 2023, 501,856 in 2024, and 526,083 in 2025. The 2025 total was 24,227 (+4.8%) above 2024. These counts measure collections recorded, not demand or people served."),
    ("text", "Group collections play a substantial role in high-volume weeks. A group log is a category of entry, not a host location; it should be reported separately from geographic hosts."),
    ("heading", "1. Growth in recorded collections"),
    ("text", "The first recorded month, April 2020, had 847 sandwiches. The program's calendar-year totals grew substantially afterward. The 2020 total covers only April 22 through December, so it is not a like-for-like full-year comparison."),
    ("table", ("Calendar year", "Sandwiches logged"), annual),
    ("text", "Monthly totals fluctuate rather than increasing every month. November 2022 is the largest complete calendar month through 2025 in these records (60,082). January 2025 and November 2023 were also high-volume months."),
    ("table", ("High-volume month", "Sandwiches logged"), months),
    ("note", "Monthly totals use collection dates within calendar months; weekly totals below use Wednesday-Tuesday windows. They are not interchangeable."),
    ("heading", "2. Locations and group contributions"),
    ("text", "The highest named host-location totals in the full 2025 calendar year were Dunwoody/PTC, Alpharetta, and East Cobb/Roswell. Intown and Sandy Springs have more than one host-name spelling in the records; those variants are combined below."),
    ("table", ("Named host location, 2025", "Sandwiches logged"), locations),
    ("text", "Entries filed under the separate 'Groups' host-name category totaled 186,327 in 2025, about 35% of the year's 526,083 sandwiches. This is not one physical distribution site and should not be ranked as a geographic host. The location totals above include any group counts recorded within those hosts' own entries."),
    ("heading", "3. Seasonal variation"),
    ("text", "Across the four full calendar years 2022-2025, November averaged 49,900 recorded sandwiches per month, followed by January at 48,273. August averaged 32,849, the lowest of the twelve months. These are four-year monthly averages, not an estimate of demand."),
    ("table", ("Month", "Average", "Month", "Average"), seasonal),
    ("text", "The pattern is not uniform: November 2024 had 31,991, far below November 2022's 60,082. A few large group entries can substantially shift monthly totals; the figures do not establish a predictable holiday or weather effect."),
    ("heading", "4. Peak sandwich collection weeks"),
    ("text", "Each row below covers Wednesday through Tuesday. 'Individual' and 'Group' refer to how sandwiches were counted within the records; their sum is the weekly total."),
    ("table", ("Completed week", "Total", "Individual", "Group"), peak_weeks),
    ("text", "The largest week, January 15-21, 2025, included a 19,706-sandwich group entry dated January 20. The November 15-21, 2023 week included a 19,414-sandwich 'Groups' entry; the full week was 27,810, not 19,414. The September 25-October 1, 2024 week totaled 24,664."),
    ("heading", "5. What the regular weeks show"),
    ("text", "The median logged Wednesday-Tuesday week rose from 8,168 in 2022 to 9,471 in 2025. Medians are less influenced than averages by unusually large group entries. The table describes weeks with a recorded total; it should not be read as a guarantee that every week had complete reporting."),
    ("table", ("Year", "Weeks with logs", "Median weekly total"), weekly_medians),
    ("text", "High-volume weeks demonstrate that more sandwiches were recorded during those periods. The collection entries alone do not tell us whether a peak was caused by a special event, a holiday campaign, an emergency response, or changes in how entries were reported."),
    ("heading", "6. Interpreting these findings"),
    ("text", "The data supports a sustained, multi-location collection operation with both recurring hosts and sizable group participation. It does not measure volunteer headcount or retention, unmet need, distribution outcomes, program costs, or the monetary value of sandwiches. Those questions require separate sources before they can be reported as facts."),
    ("heading", "Data and counting notes"),
    ("note", "Source: production sandwich_collections records, excluding soft-deleted entries. Latest completed weekly window: September 16-22, 2026; September 23-29 is still open. Calendar-year and monthly figures use collection_date. Weekly windows are Wednesday-Tuesday by collection_date, not submission date."),
    ("note", "Record total = individual_sandwiches plus the sum of group_collections counts; if the JSON group list is empty, use legacy group1_count and group2_count instead. String-encoded JSON group lists are parsed. JSON and legacy columns are never added together. Figures represent data as logged, not independently verified deliveries; category corrections can change future reports."),
]

navy = colors.HexColor("#183B51")
teal = colors.HexColor("#187481")
muted = colors.HexColor("#536674")
stripe = colors.HexColor("#F0F5F7")

# Word
word = Document()
sect = word.sections[0]
sect.top_margin = Inches(0.68)
sect.bottom_margin = Inches(0.68)
sect.left_margin = Inches(0.85)
sect.right_margin = Inches(0.85)
norm = word.styles["Normal"]
norm.font.name = "Aptos"
norm.font.size = Pt(9.6)
norm.paragraph_format.space_after = Pt(7)
for name, size, rgb in [
    ("Title", 19, RGBColor(24, 59, 81)),
    ("Subtitle", 10.5, RGBColor(24, 116, 129)),
    ("Heading 1", 12, RGBColor(24, 59, 81)),
]:
    style = word.styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.color.rgb = rgb
    style.font.bold = name != "Subtitle"
    style.paragraph_format.space_before = Pt(11)
    style.paragraph_format.space_after = Pt(5)

for item in report:
    kind = item[0]
    if kind == "heading" and item[1].startswith(("2. ", "4. ")):
        word.add_page_break()
    if kind == "title":
        word.add_heading(item[1], 0)
    elif kind == "subtitle":
        word.add_paragraph(item[1], "Subtitle")
    elif kind == "heading":
        word.add_heading(item[1], 1)
    elif kind == "table":
        labels, rows = item[1], item[2]
        table = word.add_table(rows=1, cols=len(labels))
        table.style = "Light Shading Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, value in zip(table.rows[0].cells, labels):
            cell.text = value
        for row in rows:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = value
    else:
        paragraph = word.add_paragraph(item[1])
        if kind == "note":
            paragraph.style = "Normal"
            for run in paragraph.runs:
                run.italic = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(83, 102, 116)

footer = sect.footer.paragraphs[0]
footer.text = "The Sandwich Project  |  Collection trends and peak weeks  |  September 2026"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
word.save(DOCX)

# PDF
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="ReportTitle", fontName="Helvetica-Bold", fontSize=18,
                          leading=22, textColor=navy, spaceAfter=9))
styles.add(ParagraphStyle(name="ReportSubtitle", fontName="Helvetica", fontSize=10,
                          leading=13, textColor=teal, spaceAfter=9))
styles.add(ParagraphStyle(name="ReportHeading", fontName="Helvetica-Bold", fontSize=11,
                          leading=15, textColor=navy, spaceBefore=11, spaceAfter=5))
styles.add(ParagraphStyle(name="ReportBody", fontName="Helvetica", fontSize=9,
                          leading=13.5, textColor=navy, spaceAfter=8, alignment=TA_LEFT))
styles.add(ParagraphStyle(name="ReportNote", fontName="Helvetica-Oblique", fontSize=7.7,
                          leading=11, textColor=muted, spaceAfter=7))
styles.add(ParagraphStyle(name="TableCell", fontName="Helvetica", fontSize=8,
                          leading=10, textColor=navy))
styles.add(ParagraphStyle(name="TableHead", fontName="Helvetica-Bold", fontSize=8,
                          leading=10, textColor=colors.white))

def paragraph(text, style):
    return Paragraph(escape(text), styles[style])

def make_table(labels, rows):
    # Long location labels wrap inside cells; number columns remain compact.
    if len(labels) == 4:
        widths = ([1.88, 1.6, 1.62, 1.62] if labels[0] == "Completed week"
                  else [0.9, 2.46, 0.9, 2.46])
    elif labels[0] == "Named host location, 2025":
        widths = [4.85, 1.87]
    elif len(labels) == 3:
        widths = [1.5, 2.55, 2.67]
    else:
        widths = [3.36, 3.36]
    data = [[paragraph(value, "TableHead") for value in labels]]
    data += [[paragraph(value, "TableCell") for value in row] for row in rows]
    table = Table(data, colWidths=[width * inch for width in widths], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, stripe]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.HexColor("#C6D7DF")),
    ]))
    return table

def pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D5E1E5"))
    canvas.line(0.82*inch, 0.54*inch, 7.68*inch, 0.54*inch)
    canvas.setFont("Helvetica", 7.3)
    canvas.setFillColor(muted)
    canvas.drawString(0.84*inch, 0.39*inch, "The Sandwich Project | Production collection records | September 2026")
    canvas.drawRightString(7.68*inch, 0.39*inch, str(doc.page))
    canvas.restoreState()

flow = []
for item in report:
    kind = item[0]
    if kind == "heading" and item[1].startswith(("2. ", "4. ")):
        flow.append(PageBreak())
    if kind == "table":
        flow += [make_table(item[1], item[2]), Spacer(1, 6)]
    else:
        style = {
            "title": "ReportTitle", "subtitle": "ReportSubtitle",
            "heading": "ReportHeading", "text": "ReportBody",
            "note": "ReportNote"
        }[kind]
        flow.append(paragraph(item[1], style))

pdf = SimpleDocTemplate(str(PDF), pagesize=letter, leftMargin=0.83*inch,
                        rightMargin=0.83*inch, topMargin=0.65*inch,
                        bottomMargin=0.70*inch)
pdf.build(flow, onFirstPage=pdf_footer, onLaterPages=pdf_footer)
print(DOCX)
print(PDF)